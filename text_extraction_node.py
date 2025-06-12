from PyPDF2 import PdfReader
from docx import Document
import openpyxl
import os
from typing import Dict



def extract_text_from_file(file_path: str) -> str:
    text = ""
    if file_path.endswith(".pdf"):
        try:
            reader = PdfReader(file_path)
            text = "\n".join([page.extract_text() or "" for page in reader.pages])
        except Exception as e:
            text = f"[Error reading PDF: {e}]"
    elif file_path.endswith(".docx"):
        try:
            doc = Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs])
        except Exception as e:
            text = f"[Error reading DOCX: {e}]"
    elif file_path.endswith(".xlsx"):
        try:
            wb = openpyxl.load_workbook(file_path)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join([str(cell) if cell is not None else "" for cell in row])
                    text += row_text + "\n"
        except Exception as e:
            text = f"[Error reading XLSX: {e}]"
    return text

def text_extraction_node(state: dict) -> dict:
    folder_name = state.get("folder_name")
    email_content = state.get("email_content")

    extracted_by_file = {}  # filename: text

    for file_name in os.listdir(folder_name):
        file_path = os.path.join(folder_name, file_name)
        if os.path.isfile(file_path):
            extracted_by_file[file_name] = extract_text_from_file(file_path)

    return {
        **state,
        "email_content": email_content,
        "extracted_by_file": extracted_by_file
    }
